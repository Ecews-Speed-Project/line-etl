import io
import json
import unittest
from unittest.mock import patch
from types import SimpleNamespace

from fastapi.testclient import TestClient
from openpyxl import Workbook
from app import app, save_upload


class AppTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    def test_separate_module_pages_and_assets(self):
        for path,form_id in [('/headers','review'),('/convert','convert'),('/upload','upload')]:
            response=self.client.get(path)
            self.assertEqual(response.status_code,200)
            self.assertIn('id="'+form_id+'"',response.text)
            for other in ['review','convert','upload']:
                if other!=form_id:
                    self.assertNotIn('<form id="'+other+'"',response.text)
        for path in ['/','/flat','/errors','/static/import.js','/static/import.css','/static/flat.js','/static/flat.css']:
            self.assertEqual(self.client.get(path).status_code,200)
        self.assertNotIn('<form',self.client.get('/').text)
        self.assertIn('function addCorrection',self.client.get('/static/flat.js').text)

    def test_dashboard_page_and_endpoint(self):
        self.assertEqual(self.client.get('/dashboard').status_code,200)
        self.assertEqual(self.client.get('/static/dashboard.js').status_code,200)
        with patch('routes_dashboard.summary',return_value={'linelists':[],'periods':[],'totals':{'rows':0}}):
            response=self.client.post('/api/dashboard',data={'password':'unused'})
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.json()['totals']['rows'],0)

    def test_upload_larger_than_old_limit(self):
        class Chunks:
            remaining = 101
            def read(self, size):
                if not self.remaining:
                    return b''
                self.remaining -= 1
                return b'x' * (1024 * 1024)
        class Sink:
            total = 0
            def __enter__(self):
                return self
            def __exit__(self, *args):
                pass
            def write(self, chunk):
                self.total += len(chunk)
        sink = Sink()
        save_upload(SimpleNamespace(file=Chunks()), SimpleNamespace(open=lambda mode: sink))
        self.assertEqual(sink.total, 101 * 1024 * 1024)
        with patch('web_common.MAX_SIZE', 1024):
            with self.assertRaisesRegex(ValueError, 'configured'):
                save_upload(SimpleNamespace(file=Chunks()), SimpleNamespace(open=lambda mode: Sink()))

    def test_home_and_conversion(self):
        self.assertEqual(self.client.get('/').status_code, 200)
        book = Workbook()
        book.active.append(['id', 'value'])
        book.active.append(['001', 0])
        content = io.BytesIO()
        book.save(content)
        response = self.client.post('/api/convert', files={
            'file': ('sample.xlsx', content.getvalue())})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers['x-row-count'], '1')
        self.assertEqual(response.json()['records'][0]['data'], {'id': '001', 'value': 0})

    def test_reject_invalid_files_and_origins(self):
        response = self.client.post('/api/convert', files={'file': ('bad.txt', b'bad')})
        self.assertEqual(response.status_code, 400)
        response = self.client.post('/api/convert', files={'file': ('bad.xlsx', b'bad')})
        self.assertEqual(response.status_code, 400)
        response = self.client.post('/api/upload', files={'file': ('bad.json', b'{}')},
                                    data={'password': 'unused', 'period': '2026-10-14'})
        self.assertEqual(response.status_code, 400)
        response = self.client.post('/api/convert', headers={'Origin': 'https://example.com'})
        self.assertEqual(response.status_code, 403)

    def test_mapping_routes(self):
        book = Workbook()
        book.active.append(['IP', 'IPTOutcome', 'Surname'])
        book.active.append(['ECEWS', 'Completed', 'Example'])
        buffer = io.BytesIO()
        book.save(buffer)
        content = buffer.getvalue()
        response = self.client.post('/api/headers', files={'file': ('old.xlsx', content)})
        self.assertEqual(response.status_code, 200)
        mapping = response.json()['mapping']
        mapping['TPTOutcome'] = 'IPTOutcome'
        response = self.client.post('/api/standardize', files={'file': ('old.xlsx', content)},
                                    data={'mapping': json.dumps(mapping)})
        self.assertEqual(response.status_code, 400)
        response = self.client.post('/api/standardize', files={'file': ('old.xlsx', content)},
                                    data={'mapping': json.dumps(mapping), 'allow_missing': 'true'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers['content-type'], 'application/zip')

    def test_formula_error_is_json_and_releases_file(self):
        book = Workbook()
        book.active.append(['id', 'value'])
        book.active.append(['001', '=1+1'])
        buffer = io.BytesIO()
        book.save(buffer)
        for _ in range(2):
            response = self.client.post('/api/convert', files={'file': ('formula.xlsx', buffer.getvalue())})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.headers['x-flagged-cells'], '1')
            self.assertEqual(response.json()['conversion_report']['issues'][0]['cell'], 'B2')
            self.assertIn('application/json', response.headers['content-type'])
        document = response.content
        reviewed = self.client.post('/api/review-json', files={'file': ('data.json', document)})
        self.assertEqual(reviewed.status_code, 200)
        self.assertEqual(reviewed.json()['total'], 1)
        self.assertEqual(reviewed.json()['issues'][0]['cell'], 'B2')
        blocked = self.client.post('/api/upload', files={'file': ('data.json', document)},
                                   data={'password': 'unused', 'period': '2026-10-14'})
        self.assertEqual(blocked.status_code, 400)
        self.assertIn('acknowledge', blocked.json()['detail'])
        self.assertEqual(self.client.get('/flat').status_code, 200)

    def test_review_pagination_and_flat_actions(self):
        document = {'format_version': 1, 'source_sha256': 'a'*64, 'source_file': 'old.xlsx',
                    'sheet': 'Sheet1', 'headers': ['IP'], 'records': [],
                    'conversion_report': {'issues': [{'cell': 'A2'}, {'cell': 'A3'}]}}
        response = self.client.post('/api/review-json', files={'file': ('data.json', json.dumps(document))},
                                   data={'offset': 1, 'limit': 1})
        self.assertEqual(response.json()['issues'], [{'cell': 'A3'}])
        self.assertEqual(response.json()['total'], 2)
        with patch('routes_flat.flatten', return_value={'inserted': 0, 'total': 0, 'columns': [], 'table': 'public.linelist_flat'}):
            response = self.client.post('/api/flat/build', data={'password': 'unused', 'period': '2026-10-14'})
        self.assertEqual(response.status_code, 200)

    def test_upload_requires_valid_period(self):
        document = {'format_version': 1, 'source_sha256': 'a'*64, 'source_file': 'sample.xlsx',
                    'sheet': 'Sheet1', 'headers': ['IP'], 'records': [{'source_row': 2, 'data': {'IP': 'ECEWS'}}]}
        response = self.client.post('/api/upload', files={'file': ('data.json', json.dumps(document))},
                                    data={'password': 'unused'})
        self.assertEqual(response.status_code, 422)
        response = self.client.post('/api/upload', files={'file': ('data.json', json.dumps(document))},
                                    data={'password': 'unused', 'period': '2026-02-30'})
        self.assertEqual(response.status_code, 400)
        with patch('routes_upload.upload', return_value=1) as uploader:
            response = self.client.post('/api/upload', files={'file': ('data.json', json.dumps(document))},
                                        data={'password': 'unused', 'period': '2026-10-14'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['period'], '2026-10-14')
        self.assertEqual(str(uploader.call_args.kwargs['period']), '2026-10-14')

    def test_correction_endpoint_validation(self):
        payload={'password':'unused','source_table':'linelist_records','source_sha256':'a'*64,
                 'sheet':'Sheet1','source_row':2,'heading':'LastWeight','value':'65.5'}
        with patch('routes_flat.correct',return_value={'saved':True,'heading':'LastWeight','value':'65.5'}):
            response=self.client.post('/api/correct-type',data=payload)
        self.assertEqual(response.status_code,200)
        self.assertTrue(response.json()['saved'])
        with patch('routes_flat.correct',side_effect=ValueError('Enter a valid numeric value')):
            response=self.client.post('/api/correct-type',data=payload)
        self.assertEqual(response.status_code,400)
        self.assertIn('valid numeric',response.json()['detail'])


if __name__ == '__main__':
    unittest.main()
