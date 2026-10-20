import json,tempfile,unittest,zipfile
from pathlib import Path
from datetime import datetime,date,time
from decimal import Decimal
from xml.etree import ElementTree as ET
from openpyxl import Workbook,load_workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
from header_mapping import TARGET,propose,standardize
from column_types import DATES

class DateExportTests(unittest.TestCase):
    def test_every_date_column_exports_as_short_date_without_time(self):
        with tempfile.TemporaryDirectory() as folder:
            folder=Path(folder);source=folder/'source.xlsx';book=Workbook();book.epoch=CALENDAR_MAC_1904
            headers=sorted(DATES)+['PepID','Address'];book.active.append(headers)
            vals=[datetime(2026,10,19,13,14,15) if n%2 else '2026-10-19T22:23:24+01:00' for n,_ in enumerate(sorted(DATES))]+['001','2026-10-19T10:00:00']
            book.active.append(vals);book.save(source)
            standardize(source,folder,propose(headers)['mapping'],allow_missing=True)
            data=json.loads((folder/'linelist.json').read_text())['records'][0]['data']
            exported=load_workbook(folder/'current_linelist.xlsx');ws=exported.active
            refs=set()
            for heading in DATES:
                cell=ws.cell(2,TARGET.index(heading)+1);refs.add(cell.coordinate)
                self.assertEqual(cell.number_format,'dd/mm/yyyy')
                self.assertEqual(cell.value,datetime(2026,10,19))
                self.assertEqual(data[heading],'2026-10-19')
            self.assertEqual(data['PepID'],'001');self.assertEqual(data['Address'],'2026-10-19T10:00:00')
            exported.close()
            with zipfile.ZipFile(folder/'current_linelist.xlsx') as archive:
                root=ET.fromstring(archive.read('xl/worksheets/sheet1.xml'));ns={'x':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
                for cell in root.findall('.//x:c',ns):
                    if cell.attrib['r'] in refs:
                        value=Decimal(cell.find('x:v',ns).text);self.assertEqual(value,value.to_integral_value())
    def test_missing_invalid_dates_and_legacy_alias(self):
        with tempfile.TemporaryDirectory() as folder:
            folder=Path(folder);source=folder/'source.xlsx';book=Workbook();book.active.append(['First_INH_Pickupdate','ARTStartDate','DOB']);book.active.append(['2026-10-19 17:20:30','invalid date',None]);book.save(source)
            standardize(source,folder,propose(['First_INH_Pickupdate','ARTStartDate','DOB'])['mapping'],allow_missing=True)
            out=load_workbook(folder/'current_linelist.xlsx');ws=out.active
            self.assertEqual(ws.cell(2,TARGET.index('TPT_Startdate')+1).value,datetime(2026,10,19))
            self.assertEqual(ws.cell(2,TARGET.index('ARTStartDate')+1).value,'invalid date')
            self.assertIsNone(ws.cell(2,TARGET.index('DOB')+1).value)
            out.close()
