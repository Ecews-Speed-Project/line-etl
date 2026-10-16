import json
from pathlib import Path
import tempfile
import unittest
import zipfile

from openpyxl import Workbook, load_workbook
from header_mapping import TARGET, propose, standardize


class MappingTests(unittest.TestCase):
    def test_ipt_inh_priority_aliases(self):
        aliases={'Is_Patient_Eligible_For_TPT':'Is_Patient_Eligible_For_IPT',
                 'TPT_Startdate':'First_INH_Pickupdate','TPT_Completion_date':'Last_INH_Pickupdate'}
        headers=list(aliases)+list(aliases.values())
        report=propose(headers)
        for target,source in aliases.items():
            self.assertEqual(report['mapping'][target],source)
            self.assertEqual(propose([target])['mapping'][target],target)
        with tempfile.TemporaryDirectory() as folder:
            folder=Path(folder);source=folder/'old.xlsx'
            book=Workbook();book.active.append(headers)
            book.active.append(['No','2025-01-01','2025-02-01','Yes','2026-01-01','2026-02-01'])
            book.save(source)
            result=standardize(source,folder,report['mapping'],allow_missing=True,include_excel=False)
            with zipfile.ZipFile(result) as archive:data=json.loads(archive.read('linelist.json'))['records'][0]['data']
            self.assertEqual(data['Is_Patient_Eligible_For_TPT'],'Yes')
            self.assertEqual(data['TPT_Startdate'],'2026-01-01')
            self.assertEqual(data['TPT_Completion_date'],'2026-02-01')
    def test_pbs_recaptured_automatic_alias(self):
        report=propose(['PBS_Recaptured'])
        self.assertEqual(report['mapping']['recapture'],'PBS_Recaptured')
        self.assertEqual(report['statuses']['recapture'],'automatic alias')
        self.assertNotIn('PBS_Recaptured',report['unmapped_source'])
        self.assertEqual(propose(['PBS_Recaptured','recapture'])['mapping']['recapture'],'PBS_Recaptured')
        self.assertEqual(propose(['recapture'])['mapping']['recapture'],'recapture')
        self.assertIsNone(propose(['PBS RECAPTURED','PBS Recaptured'])['mapping']['recapture'])
        with tempfile.TemporaryDirectory() as folder:
            folder=Path(folder);source=folder/'old.xlsx'
            book=Workbook();book.active.append(['PBS_Recaptured','recapture']);book.active.append([False,True]);book.save(source)
            result=standardize(source,folder,report['mapping'],allow_missing=True,include_excel=False)
            with zipfile.ZipFile(result) as archive:
                self.assertIs(json.loads(archive.read('linelist.json'))['records'][0]['data']['recapture'],False)
    def test_personal_fields_order_and_values(self):
        from column_types import TYPES
        fields=['Surname','Firstname','Educationallevel','MaritalStatus','JobStatus','PhoneNo','Address']
        start=TARGET.index('Date_Transfered_In')+1
        self.assertEqual(TARGET[start:start+7],fields)
        self.assertTrue(all(TYPES[h]=='text' for h in fields))
        with tempfile.TemporaryDirectory() as folder:
            folder=Path(folder);source=folder/'old.xlsx'
            values=['Example','Test','Secondary','Single','Employed','08001234567','Sample address']
            book=Workbook();book.active.append(fields);book.active.append(values);book.save(source)
            result=standardize(source,folder,propose(fields)['mapping'],allow_missing=True,include_excel=False)
            with zipfile.ZipFile(result) as archive:
                data=json.loads(archive.read('linelist.json'))['records'][0]['data']
            self.assertEqual([data[h] for h in fields],values)
    def test_streamed_json_only_mapping_and_progress(self):
        with tempfile.TemporaryDirectory() as folder:
            folder=Path(folder);source=folder/'old.xlsx'
            book=Workbook();book.active.append(['IP','State','Extra'])
            book.active.append(['ECEWS','AKWA IBOM','keep'])
            book.save(source)
            mapping=propose(['IP','State','Extra'])['mapping']
            events=[]
            result=standardize(source,folder,mapping,allow_missing=True,include_excel=False,progress=events.append)
            with zipfile.ZipFile(result) as archive:
                self.assertNotIn('current_linelist.xlsx',archive.namelist())
                doc=json.loads(archive.read('linelist.json'))
                original=json.loads(archive.read('original_data.json'))
                self.assertEqual(doc['records'][0]['data']['State'],'Akwa ibom')
                self.assertEqual(original['records'][0]['data']['Extra'],'keep')
                self.assertEqual(doc['mapping_report']['rows'],1)
            self.assertEqual(events[-1]['state'],'completed')
            self.assertIn('packaging',[e['phase'] for e in events])
    def test_conservative_suggestions_and_ambiguity(self):
        report = propose(['IP', 'IPTOutcome', 'Other_Test\\_(TB-LAM_LF-LAM_etc)', 'First_INH_Pickupdate', 'Last_INH_Pickupdate'])
        self.assertEqual(report['mapping']['IP'], 'IP')
        self.assertIsNone(report['mapping']['TPTOutcome'])
        self.assertEqual(report['suggestions']['TPTOutcome'], 'IPTOutcome')
        self.assertEqual(report['mapping']['TPT_Startdate'],'First_INH_Pickupdate')
        self.assertEqual(report['mapping']['Other_Test_(TB-LAM_LF-LAM_etc)'], 'Other_Test\\_(TB-LAM_LF-LAM_etc)')
        report = propose(['PATIENT ID', 'patient-id'])
        self.assertIsNone(report['mapping']['patient_id'])
        self.assertEqual(report['statuses']['patient_id'], 'ambiguous')

    def test_export_and_flags(self):
        with tempfile.TemporaryDirectory() as folder:
            folder = Path(folder)
            source = folder / 'old.xlsx'
            book = Workbook()
            book.active.append(['IP', 'IPTOutcome', 'LegacyExtra'])
            book.active.append(['ECEWS', 'Completed', 'Example'])
            book.save(source)
            mapping = propose(['IP', 'IPTOutcome', 'LegacyExtra'])['mapping']
            mapping['TPTOutcome'] = 'IPTOutcome'
            with self.assertRaisesRegex(ValueError, 'missing'):
                standardize(source, folder, mapping)
            result = standardize(source, folder, mapping, allow_missing=True)
            with zipfile.ZipFile(result) as archive:
                doc = json.loads(archive.read('linelist.json'))
                report = json.loads(archive.read('mapping_report.json'))
                original = json.loads(archive.read('original_data.json'))
            self.assertEqual(doc['headers'], TARGET)
            self.assertEqual(doc['records'][0]['data']['TPTOutcome'], 'Completed')
            self.assertIsNone(doc['records'][0]['data']['TPT_Startdate'])
            self.assertIn('LegacyExtra', report['unmapped_source'])
            self.assertEqual(original['records'][0]['data']['LegacyExtra'], 'Example')
            sheet = load_workbook(folder / 'current_linelist.xlsx', read_only=True)
            self.assertEqual(list(next(sheet.active.values)), TARGET)
            sheet.close()
            mapping['TPTOutcomeDate'] = 'IPTOutcome'
            with self.assertRaisesRegex(ValueError, 'only one'):
                standardize(source, folder, mapping, allow_missing=True)


if __name__ == '__main__':
    unittest.main()
