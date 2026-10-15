"""Convert an Excel worksheet to a provenance-aware JSON document."""
import argparse
from datetime import date, datetime, time
import hashlib
import json
from pathlib import Path
import os
import tempfile
from contextlib import ExitStack

from openpyxl import load_workbook


def normalize_state(heading, value):
    key = ''.join(c for c in heading.casefold() if c.isalnum())
    if key in ('state', 'stateofresidence') and isinstance(value, str):
        return value.strip().capitalize()
    return value


def convert(source, destination, sheet=None, header_row=1, report=None, progress=None):
    source, destination = Path(source), Path(destination)
    if source.resolve() == destination.resolve():
        raise ValueError('Input and output must be different files')
    if header_row < 1:
        raise ValueError('Header row must be positive')
    def notify(**event):
        if progress:
            progress(event)
    hasher = hashlib.sha256()
    size = source.stat().st_size
    done = 0
    notify(phase='hashing', bytes_done=0, bytes_total=size)
    with source.open('rb') as handle:
        while chunk := handle.read(1024 * 1024):
            hasher.update(chunk)
            done += len(chunk)
            notify(phase='hashing', bytes_done=done, bytes_total=size)
    digest = hasher.hexdigest()
    notify(phase='opening_workbook')
    resources = ExitStack()
    try:
        workbook_file = resources.enter_context(source.open('rb'))
        cached_file = resources.enter_context(source.open('rb'))
        workbook = load_workbook(workbook_file, read_only=True, data_only=False)
        resources.callback(workbook.close)
        cached_workbook = load_workbook(cached_file, read_only=True, data_only=True)
        resources.callback(cached_workbook.close)
    except Exception:
        resources.close()
        raise
    temp = None
    try:
        ws = workbook[sheet] if sheet else workbook.worksheets[0]
        cached_ws = cached_workbook[ws.title]
        rows = ws.iter_rows(min_row=header_row)
        cached_rows = cached_ws.iter_rows(min_row=header_row + 1)
        first = next(rows, ())
        headers = [str(c.value).strip() if c.value is not None else '' for c in first]
        while headers and not headers[-1]:
            headers.pop()
        if not headers or any(not h for h in headers) or len(set(headers)) != len(headers):
            raise ValueError('Headers must be nonempty and unique; select the correct --header-row')
        destination.parent.mkdir(parents=True, exist_ok=True)
        count = 0
        issues = []
        formula_count = 0
        total_rows = max(0, ws.max_row-header_row) if ws.max_row else None
        scanned = 0
        notify(phase='converting', processed=0, records=0, total_rows=total_rows)
        with tempfile.NamedTemporaryFile(mode='w', encoding='utf-8', dir=destination.parent,
                                         delete=False, suffix='.tmp') as output:
            temp = output.name
            metadata = dict(format_version=1, source_file=source.name, source_sha256=digest,
                            sheet=ws.title, header_row=header_row, headers=headers)
            output.write(json.dumps(metadata, ensure_ascii=False)[:-1] + ', "records": [\n')
            for row_number, (cells, cached_cells) in enumerate(zip(rows, cached_rows), header_row + 1):
                scanned = row_number-header_row
                if scanned % 500 == 0:
                    notify(phase='converting', processed=scanned, records=count, total_rows=total_rows)
                if all(c.value is None for c in cells):
                    continue
                if any(c.value is not None for c in cells[len(headers):]):
                    raise ValueError(f'Unlabelled data at row {row_number}')
                values = []
                for column, cell in enumerate(cells[:len(headers)]):
                    value = cell.value
                    cached = cached_cells[column]
                    if cell.data_type == 'f':
                        formula_count += 1
                        value = cached.value
                    issue = None
                    if cell.data_type == 'e' or cached.data_type == 'e':
                        issue = dict(kind='excel_error', error=str(cached.value if cached.data_type == 'e' else cell.value))
                        value = None
                    elif cell.data_type == 'f' and value is None:
                        issue = dict(kind='missing_formula_result')
                    if issue:
                        issues.append(dict(sheet=ws.title, cell=cell.coordinate, source_row=row_number,
                                           heading=headers[column], replacement=None, **issue))
                    if isinstance(value, (datetime, date, time)):
                        value = value.isoformat()
                    value = normalize_state(headers[column], value)
                    values.append(value)
                record = dict(source_row=row_number, data=dict(zip(headers, values)))
                if count:
                    output.write(',\n')
                output.write(json.dumps(record, ensure_ascii=False, allow_nan=False))
                count += 1
            summary = dict(formula_cells=formula_count, flagged_cells=len(issues), issues=issues,
                           note='Formula results use values saved by Excel, which may be stale. Excel errors and formulas without saved results become null; no formulas are recalculated.')
            output.write('\n], "conversion_report": ' + json.dumps(summary, ensure_ascii=False) + '}\n')
        notify(phase='finalizing', processed=scanned, records=count, total_rows=total_rows)
        os.replace(temp, destination)
        temp = None
        if report is not None:
            report.update(summary)
        return count
    finally:
        if 'rows' in locals():
            rows.close()
        if 'cached_rows' in locals():
            cached_rows.close()
        resources.close()
        if temp:
            Path(temp).unlink(missing_ok=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', type=Path)
    parser.add_argument('destination', type=Path)
    parser.add_argument('--sheet')
    parser.add_argument('--header-row', type=int, default=1)
    args = parser.parse_args()
    print(f'Exported {convert(args.source, args.destination, args.sheet, args.header_row):,} rows')


if __name__ == '__main__':
    main()
