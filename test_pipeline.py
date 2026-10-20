import json
from datetime import datetime
from pathlib import Path
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import Workbook
from excel_to_json import convert, normalize_state
from upload_json import upload


class ConversionTests(unittest.TestCase):
    def test_state_names_sentence_case(self):
        self.assertEqual(normalize_state('State', 'AKWA IBOM'), 'Akwa ibom')
        self.assertEqual(normalize_state('State_of_Residence', '  cRoSs RIVER  '), 'Cross river')
        self.assertIsNone(normalize_state('State', None))
        self.assertEqual(normalize_state('patient_id', 'ABC001'), 'ABC001')
        with tempfile.TemporaryDirectory() as folder:
            source, target = Path(folder)/'input.xlsx', Path(folder)/'output.json'
            book=Workbook()
            book.active.append(['State','State_of_Residence','patient_id'])
            book.active.append(['AKWA IBOM','cRoSs RIVER','ABC001'])
            book.save(source)
            convert(source,target)
            data=json.loads(target.read_text(encoding='utf-8'))['records'][0]['data']
            self.assertEqual(data,dict(State='Akwa ibom',State_of_Residence='Cross river',patient_id='ABC001'))
    def test_types_blanks_and_validation(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / 'input.xlsx'
            target = Path(folder) / 'output.json'
            book = Workbook()
            sheet = book.active
            sheet.append(['id', 'date', 'number', 'flag', 'missing'])
            sheet.append(['001', datetime(2026, 1, 2), 0, False, None])
            sheet.append([None] * 5)
            sheet.append(['002', None, 2, True, None])
            book.save(source)
            self.assertEqual(convert(source, target), 2)
            doc = json.loads(target.read_text(encoding='utf-8'))
            self.assertEqual(doc['records'][0]['data'], {
                'id': '001', 'date': '2026-01-02',
                'number': 0, 'flag': False, 'missing': None})
            self.assertEqual(doc['records'][1]['source_row'], 4)
            self.assertEqual(upload(target, dry_run=True), 2)
            doc['records'].append(doc['records'][0])
            target.write_text(json.dumps(doc), encoding='utf-8')
            with self.assertRaisesRegex(ValueError, 'duplicate'):
                upload(target, dry_run=True)

    def test_invalid_input_preserves_existing_output(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / 'input.xlsx'
            target = Path(folder) / 'output.json'
            target.write_text('original', encoding='utf-8')
            book = Workbook()
            book.active.append(['id', 'value'])
            book.active.append(['001', '=1+1'])
            book.save(source)
            convert(source, target)
            doc = json.loads(target.read_text(encoding='utf-8'))
            self.assertIsNone(doc['records'][0]['data']['value'])
            self.assertEqual(doc['conversion_report']['issues'][0]['kind'], 'missing_formula_result')
            target.write_text('original', encoding='utf-8')
            book.active.cell(1, 2, 'id')
            book.save(source)
            with self.assertRaisesRegex(ValueError, 'Headers'):
                convert(source, target)
            self.assertEqual(target.read_text(encoding='utf-8'), 'original')

    def test_cached_formula_and_excel_errors(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / 'input.xlsx'
            target = Path(folder) / 'output.json'
            book = Workbook()
            book.active.append(['id', 'calculated', 'error', 'formula_error'])
            book.active.append(['001', '=1+1', '#DIV/0!', '=1/0'])
            book.save(source)
            with zipfile.ZipFile(source) as archive:
                parts = {name: archive.read(name) for name in archive.namelist()}
            root = ET.fromstring(parts['xl/worksheets/sheet1.xml'])
            ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
            for cell in root.findall('.//s:c', ns):
                if cell.get('r') == 'B2':
                    cell.find('s:v', ns).text = '2'
                if cell.get('r') == 'D2':
                    cell.set('t', 'e')
                    cell.find('s:v', ns).text = '#DIV/0!'
            parts['xl/worksheets/sheet1.xml'] = ET.tostring(root)
            with zipfile.ZipFile(source, 'w') as archive:
                for name, content in parts.items():
                    archive.writestr(name, content)
            self.assertEqual(convert(source, target), 1)
            doc = json.loads(target.read_text(encoding='utf-8'))
            self.assertEqual(doc['records'][0]['data']['calculated'], 2)
            self.assertIsNone(doc['records'][0]['data']['error'])
            self.assertIsNone(doc['records'][0]['data']['formula_error'])
            self.assertEqual(doc['conversion_report']['flagged_cells'], 2)
            self.assertTrue(all(i['kind'] == 'excel_error' for i in doc['conversion_report']['issues']))
            self.assertEqual(upload(target, dry_run=True), 1)


if __name__ == '__main__':
    unittest.main()
