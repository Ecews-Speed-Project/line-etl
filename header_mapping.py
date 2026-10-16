"""Conservative header matching against the current linelist schema."""
import hashlib
import json
from pathlib import Path
import re
import zipfile

from openpyxl import Workbook, load_workbook
from excel_to_json import convert, normalize_state
from json_compression import compress_json

TARGET = json.loads((Path(__file__).parent / 'current_headers.json').read_text(encoding='utf-8'))
PRIORITY_ALIASES={
    'recapture':'PBS_Recaptured',
    'Is_Patient_Eligible_For_TPT':'Is_Patient_Eligible_For_IPT',
    'TPT_Startdate':'First_INH_Pickupdate',
    'TPT_Completion_date':'Last_INH_Pickupdate',
}
SUGGESTIONS = {
    'TB_Screening_Date': 'IPT_Screening_Date',
    'TPTOutcome': 'IPTOutcome',
    'ReasonforstoppingTPT': 'ReasonforstoppingIPT',
    'TPTOutcomeDate': 'IPTOutcomeDate',
}


def normalized(value):
    return re.sub(r'[^a-z0-9]', '', value.casefold())


def inspect_headers(source, sheet=None, header_row=1):
    if header_row < 1:
        raise ValueError('Header row must be positive')
    workbook_file = Path(source).open('rb')
    try:
        workbook = load_workbook(workbook_file, read_only=True)
    except Exception:
        workbook_file.close()
        raise
    try:
        ws = workbook[sheet] if sheet else workbook.worksheets[0]
        row = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ())
        headers = [str(v).strip() if v is not None else '' for v in row]
        while headers and not headers[-1]:
            headers.pop()
        if not headers or any(not h for h in headers) or len(set(headers)) != len(headers):
            raise ValueError('Source headings must be nonempty and unique')
        return headers
    finally:
        workbook.close()
        workbook_file.close()


def propose(headers):
    mappings, statuses = {}, {}
    for target in TARGET:
        candidates = [h for h in headers if normalized(h) == normalized(target)]
        alias=PRIORITY_ALIASES.get(target)
        aliases=[h for h in headers if normalized(h)==normalized(alias)] if alias else []
        if aliases:
            preferred=alias if alias in aliases else aliases[0] if len(aliases)==1 else None
            mappings[target]=preferred
            statuses[target]='automatic alias' if preferred else 'ambiguous'
        elif target in headers:
            mappings[target], statuses[target] = target, 'exact'
        elif len(candidates) == 1:
            mappings[target], statuses[target] = candidates[0], 'formatting match'
        else:
            mappings[target], statuses[target] = None, 'missing' if not candidates else 'ambiguous'
    used = {h for h in mappings.values() if h}
    return dict(target_headers=TARGET, source_headers=headers, mapping=mappings, statuses=statuses,
                unmapped_source=[h for h in headers if h not in used],
                suggestions={t: s for t, s in SUGGESTIONS.items() if not mappings[t] and s in headers})


def standardize(source, folder, mapping, sheet=None, header_row=1, allow_missing=False,
                include_excel=True,progress=None):
    notify=progress or (lambda event:None)
    if not isinstance(mapping, dict) or set(mapping) != set(TARGET):
        raise ValueError('Mapping must contain every current heading')
    selected=[v for v in mapping.values() if v is not None]
    if len(selected)!=len(set(selected)):
        raise ValueError('Each source heading can map to only one target heading')
    missing=[t for t in TARGET if mapping[t] is None]
    if missing and not allow_missing:
        raise ValueError('Resolve missing headings or explicitly allow empty columns')
    report={}
    book=None
    worksheet=None
    states={'State','State_of_Residence'}
    def configure(metadata):
        nonlocal book,worksheet
        headers=metadata['headers']
        if any(v not in headers for v in selected):
            raise ValueError('A selected source heading is missing from this workbook')
        report.update(mapping=mapping,missing_target=missing,
                      unmapped_source=[h for h in headers if h not in selected],
                      note='Unmapped values are preserved in original_data.json. Empty target columns contain null.')
        if include_excel:
            book=Workbook(write_only=True)
            worksheet=book.create_sheet('Current Linelist')
            worksheet.append(TARGET)
        original_hash=metadata['source_sha256']
        return dict(metadata,headers=TARGET,original_source_sha256=original_hash,
                    source_sha256=hashlib.sha256(json.dumps([original_hash,mapping,TARGET],sort_keys=True).encode()).hexdigest())
    def transform(record):
        old=record['data']
        mapped={t:old[mapping[t]] if mapping[t] else None for t in TARGET}
        for heading in states:mapped[heading]=normalize_state(heading,mapped[heading])
        if worksheet is not None:
            from openpyxl.cell import WriteOnlyCell
            values=[]
            for value in mapped.values():
                if isinstance(value,str) and value.startswith('='):
                    cell=WriteOnlyCell(worksheet,value=value);cell.data_type='s';values.append(cell)
                else:values.append(value)
            worksheet.append(values)
        return dict(record,data=mapped)
    def finish(summary,count):
        report.update(rows=count,conversion_report=summary)
        return {'mapping_report':report}
    def track(event):
        notify(dict(event,total=event.get('total_rows'),processed=event.get('processed',0)))
    try:
        count=convert(source,folder/'linelist.json',sheet,header_row,
                      metadata_transform=configure,record_transform=transform,
                      original_destination=folder/'original_data.json',final_metadata=finish,progress=track)
        if book is not None:
            notify(dict(phase='writing_excel',processed=count,total=count))
            book.save(folder/'current_linelist.xlsx')
    finally:
        if book is not None:book.close()
    compress_json(folder/'linelist.json',folder/'linelist.json.gz',progress=notify)
    (folder/'mapping_report.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    (folder/'cell_errors.json').write_text(json.dumps(report['conversion_report'],indent=2),encoding='utf-8')
    names=['linelist.json','linelist.json.gz','mapping_report.json','cell_errors.json','original_data.json']
    if include_excel:names.insert(0,'current_linelist.xlsx')
    result=folder/'converted_linelist.zip'
    with zipfile.ZipFile(result,'w',zipfile.ZIP_DEFLATED,compresslevel=1) as archive:
        for number,name in enumerate(names):
            notify(dict(phase='packaging',files_done=number,files_total=len(names)))
            archive.write(folder/name,name,compress_type=zipfile.ZIP_STORED if name.endswith(('.gz','.xlsx')) else zipfile.ZIP_DEFLATED)
    notify(dict(state='completed',phase='completed',processed=count,total=count))
    return result
